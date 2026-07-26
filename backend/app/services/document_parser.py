from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


def parse_document(path: Path, file_type: str) -> list[dict]:
    suffix = path.suffix.lower()
    if file_type == "application/pdf" or suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pptx":
        return parse_pptx(path)
    return parse_text(path)


def parse_pdf(path: Path) -> list[dict]:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        pages.append({"page": index, "text": page.extract_text() or ""})
    return pages


def parse_docx(path: Path) -> list[dict]:
    document = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return [{"page": None, "text": text}]


def parse_pptx(path: Path) -> list[dict]:
    presentation = Presentation(str(path))
    pages = []
    for index, slide in enumerate(presentation.slides, start=1):
        chunks = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                chunks.append(shape.text)
        pages.append({"page": index, "text": "\n".join(chunks)})
    return pages


def parse_text(path: Path) -> list[dict]:
    return [{"page": None, "text": path.read_text(encoding="utf-8", errors="ignore")}]
